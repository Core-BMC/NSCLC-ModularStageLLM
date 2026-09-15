from pathlib import Path
from openpyxl import Workbook, load_workbook
PROVENANCE = "These examples were created solely to demonstrate the input format. They are entirely fictional and are not derived from patient records."
def generate(directory):
    directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    headers=['hospital_id','Pathology','Chest CT','Brain MR','PET','EBUS','neck biopsy','Bone scan','Abdomen&Pelvis CT','Adrenal CT','cT','cN','cM','cStage']
    files=[]
    for edition,size,site in [('8th','1.3','right middle lobe'),('9th','1.6','right lower lobe')]:
        row=['SYNTHETIC_AJCC_'+edition.upper()+'_001','Fictional example: biopsy identifies invasive adenocarcinoma.',f'Fictional example: solitary solid {size} cm nodule in the {site}. No invasion of adjacent structures. No enlarged regional nodes.','Fictional example: no brain metastasis.','Fictional example: uptake confined to the lung nodule; no nodal or distant metastatic findings.',None,None,None,None,None,'T1b','N0','M0','IA2']
        wb=Workbook();ws=wb.active;ws.title='Synthetic input';ws.append(headers);ws.append(row);ws.freeze_panes='A2'
        from openpyxl.styles import Font,Alignment,PatternFill
        for cell in ws[1]:
            cell.font=Font(name='Arial',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='245778')
        for cell in ws[2]:
            cell.font=Font(name='Arial',size=11);cell.alignment=Alignment(wrap_text=True,vertical='top')
        for col in range(1,15):ws.column_dimensions[ws.cell(1,col).column_letter].width=42 if 2<=col<=5 else 28
        ws.row_dimensions[2].height=130
        wb.properties.creator='Hwon Heo';wb.properties.description=PROVENANCE
        f=directory/('synthetic_single_excel_ajcc_'+edition+'.xlsx');wb.save(f)
        assert list(load_workbook(f,data_only=True).active.values)==[tuple(headers),tuple(row)]
        files.append(f)
    return files
if __name__=='__main__':
    generate(Path(__file__).resolve().parent)
