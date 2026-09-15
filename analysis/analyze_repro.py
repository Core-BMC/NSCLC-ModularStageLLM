"""AR-6 reproducibility analysis: run-to-run label agreement + accuracy mean±SD.

r1 is REUSED from the existing full try1 runs (output/json/*.json) restricted to the
50-case subset via the seed-42 mapping (identical to make_repro_subset.py); r2, r3, ...
come from output/repro/*.json. Case predictions are batch-independent, so try1's output
for a subset case is a valid first replicate.

Run: python3 analyze_repro.py   (after r2/r3 collected in output/repro/)
Outputs -> analysis/results/repro_ar6.md, repro_ar6.csv
"""
from __future__ import annotations

import glob
import json
import os
import random
import re
from collections import defaultdict
from itertools import combinations
from statistics import mean, stdev

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
# Reference-standard labels (case identifier + cT/cN/cM/stage per edition).
# Not distributed with the repository: the file is held under institutional
# data governance. Override with the REFERENCE_XLSX environment variable.
REFERENCE_XLSX = os.environ.get("REFERENCE_XLSX", "reference_labels.xlsx")
REPO = os.path.dirname(HERE)
GT_FULL = os.path.join(REPO, "input", REFERENCE_XLSX)
FULL_DIR = os.path.join(REPO, "output", "json")     # r1 source (full try1)
REPRO_DIR = os.path.join(REPO, "output", "repro")   # r2, r3, ...
OUT = os.path.join(HERE, "results")
os.makedirs(OUT, exist_ok=True)

SEED, N_SUBSET = 42, 50
RUN_RE = re.compile(r"(llama3-70b|phi4|gpt4o)-(mma|base)-ajcc([89])", re.I)
REP_RE = re.compile(r"repro-r(\d+)", re.I)
MODEL = {"llama3-70b": "LLaMA 3.3 70B", "phi4": "Phi-4 14B", "gpt4o": "GPT-4o"}
MODE = {"mma": "MAA", "base": "baseline"}
AXES = [("T", "T"), ("N", "N"), ("M", "M"), ("S", "Stage")]


def load_gt_full():
    wb = openpyxl.load_workbook(GT_FULL, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    hdr = [str(h) if h is not None else "" for h in next(rows)]
    ix = {n: i for i, n in enumerate(hdr)}
    out = []
    for r in rows:
        if r is None or all(c is None for c in r):
            continue
        g = lambda c: ("" if (c not in ix or r[ix[c]] is None) else str(r[ix[c]]).strip())
        out.append({"8": {"T": g("8th-cT"), "N": g("8th-cN"), "M": g("8th-cM"), "S": g("8th-cStage")},
                    "9": {"T": g("9th-cT"), "N": g("9th-cN"), "M": g("9th-cM"), "S": g("9th-cStage")}})
    return out


def cfg_key(name):
    m = RUN_RE.search(name)
    return (m.group(1).lower(), m.group(2).lower(), m.group(3)) if m else None


def labels_by_id(fp):
    d = json.load(open(fp))
    return {rec.get("id"): {a: (rec.get("aiTnm") or {}).get(k, "") or "" for a, k in AXES} for rec in d}


def main():
    gt = load_gt_full()
    idx = sorted(random.Random(SEED).sample(range(len(gt)), N_SUBSET))  # subset pos -> orig row
    # r1 from full try1
    reps = defaultdict(dict)  # key -> {repeat_no: {subset_pos: {axis:label}}}
    for fp in glob.glob(os.path.join(FULL_DIR, "*.json")):
        if "repro" in os.path.basename(fp).lower():
            continue
        k = cfg_key(os.path.basename(fp))
        if not k:
            continue
        by_id = labels_by_id(fp)
        reps[k][1] = {pos: by_id.get(idx[pos] + 1, {a: "" for a, _ in AXES}) for pos in range(N_SUBSET)}
    # r2.. from repro
    for fp in glob.glob(os.path.join(REPRO_DIR, "*.json")):
        k = cfg_key(os.path.basename(fp))
        rm = REP_RE.search(os.path.basename(fp))
        if not k or not rm:
            continue
        by_id = labels_by_id(fp)  # subset run: id = pos+1
        reps[k][int(rm.group(1))] = {pos: by_id.get(pos + 1, {a: "" for a, _ in AXES}) for pos in range(N_SUBSET)}

    if not reps:
        print("No runs found."); return

    rows_csv = ["model,mode,edition,axis,n_repeats,acc_mean,acc_sd,unanimous_rate,pairwise_agreement,"
                "consistently_correct,correctness_stable"]
    md = ["# AR-6 reproducibility — 50-case subset (seed 42); r1=full try1 reused, r2+=repeats\n",
          "acc = per-repeat accuracy vs GT (mean±SD). unanimous = % subset cases with identical "
          "label across ALL repeats. pairwise = mean pairwise label agreement. consistently_correct = "
          "% valid cases correct in ALL repeats. correctness_stable = % valid cases with identical "
          "correct/incorrect status across ALL repeats.\n",
          "\n| Model | Mode | Ed | Axis | repeats | acc mean±SD | unanimous | pairwise | consist. correct | correctness stable |",
          "|---|---|---|---|---|---|---|---|---|---|"]
    for (model, mode, ed), rr in sorted(reps.items()):
        rlist = sorted(rr)
        if len(rlist) < 2:
            md.append(f"| {MODEL[model]} | {MODE[mode]} | {ed}th | (only r{rlist} present — run r2/r3) |||||")
            continue
        for a, _ in AXES:
            accs = []
            for r in rlist:
                ok = tot = 0
                for pos in range(N_SUBSET):
                    g = gt[idx[pos]][ed][a]
                    if g == "":
                        continue
                    tot += 1
                    ok += (rr[r][pos][a] == g)
                accs.append(100 * ok / tot if tot else float("nan"))
            am, asd = mean(accs), (stdev(accs) if len(accs) > 1 else 0.0)  # sample SD (n-1)
            unan = pw = npw = 0
            cc = cs = valid = 0  # consistently-correct / correctness-stable, over valid (GT != "") cases
            for pos in range(N_SUBSET):
                labs = [rr[r][pos][a] for r in rlist]
                unan += all(x == labs[0] for x in labs)
                for x, y in combinations(labs, 2):
                    pw += (x == y); npw += 1
                g = gt[idx[pos]][ed][a]
                if g == "":
                    continue
                valid += 1
                corr = [lab == g for lab in labs]
                cc += all(corr)                       # correct in ALL repeats
                cs += (all(corr) or not any(corr))    # same correct/incorrect status in ALL repeats
            ur = 100 * unan / N_SUBSET
            pr = 100 * pw / npw if npw else float("nan")
            ccr = 100 * cc / valid if valid else float("nan")
            csr = 100 * cs / valid if valid else float("nan")
            rows_csv.append(f"{MODEL[model]},{MODE[mode]},{ed}th,{a},{len(rlist)},{am:.17g},{asd:.17g},{ur:.1f},{pr:.1f},{ccr:.1f},{csr:.1f}")
            md.append(f"| {MODEL[model]} | {MODE[mode]} | {ed}th | {a} | {len(rlist)} | {am:.1f}±{asd:.2f} | {ur:.1f}% | {pr:.1f}% | {ccr:.1f}% | {csr:.1f}% |")

    open(os.path.join(OUT, "repro_ar6.csv"), "w").write("\n".join(rows_csv) + "\n")
    open(os.path.join(OUT, "repro_ar6.md"), "w").write("\n".join(md) + "\n")
    print("\n".join(md))
    print(f"\nWrote {OUT}/repro_ar6.md, repro_ar6.csv")


if __name__ == "__main__":
    main()
