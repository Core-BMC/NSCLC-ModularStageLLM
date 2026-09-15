"""AR-6 reproducibility: build a fixed 50-case random subset from the 495-case set.

Deterministic (seed=42). Preserves all columns and sheet name 'data' so the existing
per-run configs run unchanged (just point --i at the subset). Output is gitignored (PHI).

Run: python3 make_repro_subset.py
"""
from __future__ import annotations

import os
import random

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
# Reference-standard labels (case identifier + cT/cN/cM/stage per edition).
# Not distributed with the repository: the file is held under institutional
# data governance. Override with the REFERENCE_XLSX environment variable.
REFERENCE_XLSX = os.environ.get("REFERENCE_XLSX", "reference_labels.xlsx")
REPO = os.path.dirname(HERE)
SRC = os.path.join(REPO, "input", REFERENCE_XLSX)
DST = os.path.join(REPO, "input", "reference_labels_repro50.xlsx")
N_SUBSET = 50
SEED = 42


def main() -> None:
    wb = openpyxl.load_workbook(SRC)
    ws = wb["data"]
    rows = list(ws.iter_rows(values_only=True))
    header, data = rows[0], [r for r in rows[1:] if r and not all(c is None for c in r)]
    idx = sorted(random.Random(SEED).sample(range(len(data)), N_SUBSET))  # keep original order
    picked = [data[i] for i in idx]

    out = openpyxl.Workbook()
    sh = out.active
    sh.title = "data"
    sh.append(list(header))
    for r in picked:
        sh.append(list(r))
    out.save(DST)

    nos = [str(data[i][0]) for i in idx]
    print(f"wrote {DST}  ({N_SUBSET} rows, seed={SEED})")
    print("sampled original 'No.':", ", ".join(nos))


if __name__ == "__main__":
    main()
