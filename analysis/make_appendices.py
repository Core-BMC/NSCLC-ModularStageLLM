"""Build Multimedia Appendices 5-8 as workbooks, from the analysis outputs.

Nothing here is retyped: every value is read from analysis/deposit/ or from
analysis/results*/, so an appendix cannot drift from the tables it supports.

  MA5  component-level paired comparison of the three configurations (72 rows)
  MA6  per-class precision, recall and F1 with class support
  MA7  parse-path audit over the stored raw model outputs
  MA8  histology sensitivity analyses (N=480 narrow, N=477 broad)

Run:  python3 analysis/make_appendices.py [OUT_DIR]
"""
from __future__ import annotations

import csv
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
DEP = os.path.join(HERE, "deposit")
RES = os.path.join(HERE, "results")
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "appendices")

ED = {"AJCC8th": "AJCC 8th", "AJCC9th": "AJCC 9th", "8th": "AJCC 8th", "9th": "AJCC 9th"}
CFG = {"baseline": "Single prompt", "decomposition-only": "Decomposition only", "MAA": "MAA"}
AX = {"T": "cT", "N": "cN", "M": "cM", "S": "Overall stage", "stage group": "Overall stage",
      "cT": "cT", "cN": "cN", "cM": "cM", "Stage": "Overall stage"}
PAIR = {("baseline", "decomposition-only"): "Decomposition",
        ("decomposition-only", "MAA"): "Voting",
        ("baseline", "MAA"): "Combined"}


def sheet(wb, title, note, header, rows, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append([note])
    ws.cell(1, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(header), 6))
    ws.row_dimensions[1].height = 60
    ws.append([])
    ws.append(header)
    for c in range(1, len(header) + 1):
        ws.cell(3, c).font = Font(bold=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A4"
    for c in range(1, len(header) + 1):
        width = max([len(str(header[c - 1]))] + [len(str(r[c - 1])) for r in rows if c <= len(r)]) + 2
        ws.column_dimensions[get_column_letter(c)].width = min(width, 46)
    return ws


def fmt_p(p: float) -> str:
    """JMIR P formatting, identical to the manuscript tables."""
    if p < .001:
        return "<.001"
    if p > .99:
        return ">.99"
    return f"{p:.3f}".lstrip("0")


# ---- MA5 -----------------------------------------------------------------
def ma5():
    rows = []
    for r in csv.DictReader(open(os.path.join(DEP, "mcnemar_contingency.csv"))):
        label = PAIR[(r["configuration_1"], r["configuration_2"])]
        if label == "Voting" and r["edition"] == "AJCC9th":
            label = "Voting + wording"
        p = float(r["p_exact_mcnemar"])
        rows.append([ED[r["edition"]], r["model"], AX[r["component"]], label,
                     CFG[r["configuration_1"]], CFG[r["configuration_2"]],
                     int(r["n_paired"]), float(r["accuracy_1_pct"]), float(r["accuracy_2_pct"]),
                     int(r["both_correct"]), int(r["only_1_correct"]),
                     int(r["only_2_correct"]), int(r["both_incorrect"]),
                     fmt_p(p), "yes" if p < .05 else ""])
    order = {"AJCC 8th": 0, "AJCC 9th": 1}
    axo = {"cT": 0, "cN": 1, "cM": 2, "Overall stage": 3}
    pairo = {"Decomposition": 0, "Voting": 1, "Voting + wording": 1, "Combined": 2}
    rows.sort(key=lambda r: (order[r[0]], r[1], axo[r[2]], pairo[r[3]]))
    note = ("Multimedia Appendix 5. Paired comparisons of the three configurations for cT, cN, cM and the "
            "overall stage group. All 72 comparisons: three configuration pairs for each of three models, "
            "two AJCC editions and four outcomes, on the identical 495 cases in each pair. "
            "b is the number of cases correct in configuration 1 only and c the number correct in "
            "configuration 2 only; exact McNemar P values were calculated from these discordant-pair "
            "counts. For the ninth-edition analyses the Voting comparison is labelled Voting + wording, "
            "because the configurations differ in both the voting setting and the instructions for "
            "equivocal findings. That comparison does not isolate the effect of voting. P values are "
            "uncorrected for multiplicity; 37 of the 72 comparisons reach P<.05. "
            "Source: analysis/deposit/mcnemar_contingency.csv.")
    wb = Workbook()
    sheet(wb, "72 paired comparisons", note,
          ["Edition", "Model", "Outcome", "Comparison", "Configuration 1", "Configuration 2",
           "n paired", "Accuracy 1 (%)", "Accuracy 2 (%)", "Both correct", "b (1 only)",
           "c (2 only)", "Both incorrect", "P (exact)", "P<.05"], rows, first=True)
    return wb, len(rows)


# ---- MA6 -----------------------------------------------------------------
def ma6():
    rows = []
    for r in csv.DictReader(open(os.path.join(RES, "per_class_f1.csv"))):
        rows.append([ED[r["edition"]], r["model"], CFG[r["mode"]], AX[r["axis"]], r["class"],
                     float(r["precision"]), float(r["recall"]), float(r["f1"]), int(r["support"])])
    note = ("Multimedia Appendix 6. Per-class precision, recall and F1 with class support, for every "
            "model, configuration, AJCC edition and TNM component. Support is the number of cases "
            "carrying that class in the reference standard (N=495 in every run). The macro-averaged F1 "
            "reported in Table 5 is the unweighted mean of the F1 column within each model x "
            "configuration x edition x component block. Each class therefore contributes equally to "
            "macro-F1, irrespective of its support: rare classes, including ninth-edition M1c1 with "
            "eight cases, receive the same weight as more frequent classes. "
            "Source: analysis/results/per_class_f1.csv.")
    wb = Workbook()
    sheet(wb, "Per-class metrics", note,
          ["Edition", "Model", "Configuration", "Component", "Class",
           "Precision (%)", "Recall (%)", "F1 (%)", "Support (n)"], rows, first=True)
    return wb, len(rows)


# ---- MA7 -----------------------------------------------------------------
def ma7():
    rows, tot, js, rg, no = [], 0, 0, 0, 0
    for r in csv.DictReader(open(os.path.join(RES, "parse_audit.csv"))):
        n = int(r["n_outputs"])
        rows.append([ED[r["edition"]], r["model"], CFG[r["configuration"]], r["component"], n,
                     int(r["json"]), int(r["regex"]), int(r["none"]),
                     float(r["json_fail_pct"]), float(r["regex_pct"]), float(r["none_pct"])])
        tot += n; js += int(r["json"]); rg += int(r["regex"]); no += int(r["none"])
    rows.append(["Total", "", "", "", tot, js, rg, no,
                 round(100 * (rg + no) / tot, 2), round(100 * rg / tot, 2), round(100 * no / tot, 2)])
    note = ("Multimedia Appendix 7. Parse-path audit. The audit applied the repository parsers to the "
            "stored model outputs and classified each parsing result as the JSON path, the "
            "regular-expression fallback, or no extracted label. The agent configurations were audited "
            "per component (cT, cN, cM); the single-prompt configuration issues one combined response "
            "per case and was audited on that response with TNMOutputParser, the component parsers "
            "never being used there. Outputs with no extracted label are recorded as a (no label) entry "
            "in the deposited confusion matrices and scored as incorrect. These parsing failures are "
            "distinct from an indeterminate category (cTx, cNx, cMx), which is a judgement returned by "
            "the model rather than a parse failure. The audit covers the outputs that were stored; "
            "generations discarded or re-sampled during a run were not retained. The audit script "
            "imports the parser implementations directly from src/: analysis/parse_audit.py. "
            "Source: analysis/results/parse_audit.csv.")
    wb = Workbook()
    ws = sheet(wb, "Parse-path audit", note,
               ["Edition", "Model", "Configuration", "Component", "Outputs (n)",
                "JSON path", "Regex fallback", "Neither", "JSON-fail (%)",
                "Regex (%)", "Neither (%)"], rows, first=True)
    for c in range(1, 12):
        ws.cell(ws.max_row, c).font = Font(bold=True)
    return wb, len(rows) - 1


# ---- MA8 -----------------------------------------------------------------
def _t4(path, n_cases):
    """Accuracy or macro-F1 table; the macro-F1 file carries no n column."""
    rows = []
    for r in csv.DictReader(open(path)):
        val = r.get("accuracy") or r["macro_f1"]
        rows.append([ED[r["edition"]], r["model"], CFG[r["mode"]], AX[r["axis"]],
                     int(r.get("n") or n_cases), float(val), float(r["ci_lo"]), float(r["ci_hi"])])
    return rows


def _mcn_md(path):
    """The mcnemar.md tables, relabelled to match the accuracy sheets and with
    the same P formatting as the manuscript (the raw file carries a trailing
    asterisk on significant values, which nothing in the appendix defines)."""
    rows = []
    for line in open(path, encoding="utf-8"):
        if not line.startswith("|") or "---" in line:
            continue
        c = [t.strip() for t in line.strip().strip("|").split("|")]
        if len(c) < 8 or c[0] == "Model":
            continue
        raw = c[-1].strip().rstrip("*")
        p = float(raw[1:]) / 2 if raw.startswith("<") else float(raw)
        a, b = c[3].split(" vs ")
        rows.append([ED[c[1]], c[0], AX[c[2]], f"{CFG[a]} vs {CFG[b]}", int(c[4]),
                     float(c[5]), float(c[6]), int(c[7]), int(c[8]),
                     fmt_p(p), "yes" if p < .05 else ""])
    return rows


def ma8():
    note = ("Multimedia Appendix 8. Histology sensitivity analyses. Eligibility was based on the "
            "availability of a staging work-up rather than on histologic confirmation, so the cohort "
            "contains 15 patients with a small-cell component and, in addition, one large cell "
            "neuroendocrine carcinoma, one carcinoid tumour and one mucoepidermoid carcinoma. The "
            "analyses were repeated after two exclusions: the 15 patients with a small-cell component "
            "(narrow, N=480), and those 15 together with the three other confirmed non-NSCLC "
            "histologies (broad, N=477). Against the full cohort the largest change in any accuracy is "
            "0.8 percentage points under either exclusion, and the direction of the configuration "
            "differences was unchanged. Under the narrow exclusion no comparison crosses P=.05. Under "
            "the broad exclusion one of the 72 comparisons crosses the nominal threshold: Phi-4 14B, "
            "AJCC 9th edition, cT, single prompt versus decomposition-only, P=.040 to P=.052. "
            "Source: analysis/results_excl_narrow/ and analysis/results_excl_broad/.")
    wb = Workbook()
    h4 = ["Edition", "Model", "Configuration", "Component", "n", "Accuracy (%)", "95% CI low", "95% CI high"]
    h5 = ["Edition", "Model", "Configuration", "Component", "n", "Macro-F1 (%)", "95% CI low", "95% CI high"]
    hm = ["Edition", "Model", "Outcome", "Comparison", "n paired", "Accuracy 1 (%)",
          "Accuracy 2 (%)", "b (1 only)", "c (2 only)", "P (exact)", "P<.05"]
    n = 0
    first = True
    for tag, d, n_cases in (("narrow (N=480)", "results_excl_narrow", 480),
                            ("broad (N=477)", "results_excl_broad", 477)):
        base = os.path.join(HERE, d)
        for name, path, head, load in (
                (f"Accuracy {tag}", "table4.csv", h4, _t4),
                (f"Macro-F1 {tag}", "table4_f1.csv", h5, _t4),
                (f"McNemar {tag}", "mcnemar.md", hm, _mcn_md)):
            rows = load(os.path.join(base, path), n_cases) if load is _t4 else load(os.path.join(base, path))
            sheet(wb, name[:31], note, head, rows, first=first)
            first = False
            n += len(rows)
    return wb, n


def main():
    os.makedirs(OUT, exist_ok=True)
    for num, fn in ((5, ma5), (6, ma6), (7, ma7), (8, ma8)):
        wb, n = fn()
        path = os.path.join(OUT, f"Multimedia_Appendix_{num}.xlsx")
        wb.save(path)
        print(f"  Multimedia Appendix {num}: {n} rows -> {path}")


if __name__ == "__main__":
    main()
