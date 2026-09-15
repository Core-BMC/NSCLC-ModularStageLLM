"""Analysis pipeline for JMIR #101332 — Table 4, bootstrap 95% CI, McNemar, Figure 3.

Auto-discovers every run under <repo>/output/json/ and <repo>/output/ itself, so the
originally reported baseline/MAA runs and the decomposition-only ablation runs written
by config/runs/run_deconly_matrix.sh are analysed together. Override the search
locations with JSON_DIRS (colon-separated, relative to the repository root).

Aligns predictions to ground truth by JSON `id` (1..N) = xlsx data-row order. Empty or
failed predictions count as incorrect. A run whose JSON holds fewer cases than the
reference standard is reported with its own n, so a partially completed run is visible
as such rather than silently scored against the full set.

Run:  python3 analyze_results.py
Outputs -> analysis/results/{table4.md,table4.csv,mcnemar.md,confusion/*.csv,figures/*.png}
"""
from __future__ import annotations

import glob
import hashlib
import json
import os
import re
from collections import Counter, defaultdict
from statistics import mean
from typing import Dict, List, Optional, Tuple

import numpy as np
import openpyxl

# ---- paths (relative to this script -> portable Mac/sandbox) -------------
HERE = os.path.dirname(os.path.abspath(__file__))  # <repo>/analysis
# Reference-standard labels (case identifier + cT/cN/cM/stage per edition).
# Not distributed with the repository: the file is held under institutional
# data governance. Override with the REFERENCE_XLSX environment variable.
# A bare name is looked up under <repo>/input; an absolute path is used as given, so
# the label file can be kept outside the repository altogether.
REFERENCE_XLSX = os.environ.get("REFERENCE_XLSX", "reference_labels.xlsx")
REPO = os.path.dirname(HERE)
GT_XLSX = (REFERENCE_XLSX if os.path.isabs(REFERENCE_XLSX)
           else os.path.join(REPO, "input", REFERENCE_XLSX))
# Search locations for run JSON, in order. `output/json` holds the originally reported
# runs; `output` holds the ablation runs written by config/runs/run_deconly_matrix.sh.
JSON_DIRS = [os.path.join(REPO, d) for d in
             os.environ.get("JSON_DIRS", "output/json:output").split(":") if d]
OUT = os.environ.get("RESULTS_DIR") or os.path.join(HERE, "results")
OUT = OUT if os.path.isabs(OUT) else os.path.join(HERE, OUT)
FIG = os.path.join(OUT, "figures")
CONF = os.path.join(OUT, "confusion")
for d in (OUT, FIG, CONF):
    os.makedirs(d, exist_ok=True)

# Optional filters, for interim analyses while some runs are still executing:
#   MODELS=gpt4o,phi4   EDITIONS=8,9
# An unfiltered run analyses whatever is on disk, incomplete runs included.
ONLY_MODELS = {m.strip().lower() for m in os.environ.get("MODELS", "").split(",") if m.strip()}
ONLY_EDITIONS = {e.strip() for e in os.environ.get("EDITIONS", "").split(",") if e.strip()}
# Sensitivity analysis: drop these case ids from every run before scoring, e.g.
#   EXCLUDE_CASES=3,17,204   or   EXCLUDE_CASES=@path/to/ids.txt (one id per line)
def _load_excluded() -> set:
    spec = os.environ.get("EXCLUDE_CASES", "").strip()
    if not spec:
        return set()
    if spec.startswith("@"):
        with open(spec[1:]) as fh:
            spec = ",".join(l.strip() for l in fh if l.strip() and not l.startswith("#"))
    return {int(x) for x in spec.replace("\n", ",").split(",") if x.strip().isdigit()}
EXCLUDED = _load_excluded()

SEED = 42
RNG = np.random.default_rng(SEED)
N_BOOT = 1000


def cell_rng(key: str) -> "np.random.Generator":
    """A generator determined by the cell alone, not by what ran before it.

    A single module-level generator consumed sequentially makes every interval
    depend on how many cells were computed earlier, so adding a run silently
    shifts the intervals of every later cell while leaving the point estimates
    untouched. That happened between the first and second revisions. Seeding per
    cell removes the coupling: an interval is now a function of (cell, SEED,
    N_BOOT) and of nothing else, so it is reproducible in isolation and stable
    when runs are added or the analysis is restricted.
    """
    h = hashlib.sha256(f"{SEED}|{key}".encode()).digest()[:8]
    return np.random.default_rng(int.from_bytes(h, "big"))

# ---- category orders -----------------------------------------------------
T_ORDER = ["Tis", "T1mi", "T1a", "T1b", "T1c", "T2a", "T2b", "T3", "T4", "T0", "Tx"]
N_ORDER = {"8": ["N0", "N1", "N2", "N3", "Nx"],
           "9": ["N0", "N1", "N2a", "N2b", "N3", "Nx"]}
M_ORDER = {"8": ["M0", "M1a", "M1b", "M1c", "Mx"],
           "9": ["M0", "M1a", "M1b", "M1c1", "M1c2", "Mx"]}
S_ORDER = ["0", "IA1", "IA2", "IA3", "IB", "IIA", "IIB", "IIIA", "IIIB", "IIIC", "IVA", "IVB"]


def fam_N(v: str) -> str:
    return "N2" if v in ("N2", "N2a", "N2b") else v


def fam_M(v: str) -> str:
    return "M1c" if v in ("M1c", "M1c1", "M1c2") else v


# ---- ground truth --------------------------------------------------------
def load_gt() -> List[Dict[str, str]]:
    wb = openpyxl.load_workbook(GT_XLSX, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    hdr = [str(h) if h is not None else "" for h in next(rows)]
    idx = {name: i for i, name in enumerate(hdr)}
    out = []
    for r in rows:
        if r is None or all(c is None for c in r):
            continue
        g = lambda c: ("" if (c not in idx or r[idx[c]] is None) else str(r[idx[c]]).strip())
        out.append({"8": {"T": g("8th-cT"), "N": g("8th-cN"), "M": g("8th-cM"), "S": g("8th-cStage")},
                    "9": {"T": g("9th-cT"), "N": g("9th-cN"), "M": g("9th-cM"), "S": g("9th-cStage")}})
    return out


# ---- run discovery -------------------------------------------------------
RUN_RE = re.compile(r"(llama3-70b|phi4|gpt4o)-(mma|deconly|base)-ajcc([89])", re.I)
MODEL_LABEL = {"llama3-70b": "LLaMA 3.3 70B", "phi4": "Phi-4 14B", "gpt4o": "GPT-4o"}
MODE_LABEL = {"base": "baseline", "deconly": "decomposition-only", "mma": "MAA"}
MODE_ORDER = ["base", "deconly", "mma"]   # condition A, C, D
# Table rows are keyed by display label; keep them in condition order, not alphabetical.
LABEL_RANK = {MODE_LABEL[m]: i for i, m in enumerate(MODE_ORDER)}


def discover() -> List[dict]:
    runs: List[dict] = []
    seen: Dict[Tuple[str, str, str], str] = {}
    for d in JSON_DIRS:
        for fp in sorted(glob.glob(os.path.join(d, "*.json"))):
            base = os.path.basename(fp)
            if "repro" in base.lower():   # reproducibility repeats are analysed separately
                continue
            m = RUN_RE.search(base)
            if not m:
                continue
            model, mode, ed = m.group(1).lower(), m.group(2).lower(), m.group(3)
            if ONLY_MODELS and model not in ONLY_MODELS:
                continue
            if ONLY_EDITIONS and ed not in ONLY_EDITIONS:
                continue
            key = (model, mode, ed)
            if key in seen:
                print(f"  ! {base} duplicates {seen[key]} for {key} - keeping the first")
                continue
            with open(fp) as fh:
                data = json.load(fh)
            data.sort(key=lambda r: r.get("id", 0))
            if EXCLUDED:
                data = [r for r in data if r.get("id") not in EXCLUDED]
            seen[key] = base
            runs.append({"file": base, "model": model, "mode": mode,
                         "ed": ed, "data": data})
    return runs


def preds(run: dict, axis: str) -> List[str]:
    key = axis if axis != "S" else "Stage"
    return [(r.get("aiTnm") or {}).get(key if axis != "S" else "Stage", "") or "" for r in run["data"]]


def gts(gt: List[dict], run: dict, axis: str) -> List[str]:
    """Reference labels aligned to this run's cases by id (1-based row order).

    Aligning by position would be wrong whenever cases have been dropped, as in a
    sensitivity analysis, so the id carried in each record is used instead.
    """
    return [gt[rec["id"] - 1][run["ed"]][axis] if 1 <= rec.get("id", 0) <= len(gt) else ""
            for rec in run["data"]]


# ---- metrics -------------------------------------------------------------
def accuracy_ci(pred: List[str], truth: List[str], key: str = "") -> Tuple[float, float, float, int]:
    p = np.array(pred, dtype=object)
    t = np.array(truth, dtype=object)
    mask = t != ""
    p, t = p[mask], t[mask]
    n = len(t)
    correct = (p == t)
    acc = correct.mean() * 100 if n else float("nan")
    rng = cell_rng("acc|" + key) if key else RNG
    boots = []
    for _ in range(N_BOOT):
        idx = rng.integers(0, n, n)
        boots.append(correct[idx].mean() * 100)
    lo, hi = np.percentile(boots, [2.5, 97.5])
    return acc, lo, hi, n


def accuracy_lenient(pred: List[str], truth: List[str], axis: str) -> float:
    f = fam_N if axis == "N" else fam_M
    ok = tot = 0
    for p, t in zip(pred, truth):
        if t == "":
            continue
        tot += 1
        ok += (f(p) == f(t))
    return 100 * ok / tot if tot else float("nan")


def _per_class(pred, truth):
    """Per-class precision/recall/F1/support over classes present in GT (one-vs-rest)."""
    classes = sorted({t for t in truth if t != ""})
    tp = Counter(); fp = Counter(); fn = Counter(); sup = Counter()
    for p, t in zip(pred, truth):
        if t == "":
            continue
        sup[t] += 1
        if p == t:
            tp[t] += 1
        else:
            fn[t] += 1
            if p:
                fp[p] += 1
    per = {}
    for c in classes:
        prec = tp[c] / (tp[c] + fp[c]) if (tp[c] + fp[c]) else 0.0
        rec = tp[c] / (tp[c] + fn[c]) if (tp[c] + fn[c]) else 0.0
        f1 = 2 * prec * rec / (prec + rec) if (prec + rec) else 0.0
        per[c] = (prec * 100, rec * 100, f1 * 100, sup[c])
    return per


def macro_f1(pred, truth) -> float:
    per = _per_class(pred, truth)
    return mean(v[2] for v in per.values()) if per else float("nan")


def macro_f1_ci(pred, truth, key: str = "") -> Tuple[float, float, float]:
    p = np.array(pred, dtype=object); t = np.array(truth, dtype=object)
    mask = t != ""; p, t = p[mask], t[mask]
    n = len(t)
    val = macro_f1(list(p), list(t))
    rng = cell_rng("f1|" + key) if key else RNG
    boots = []
    for _ in range(N_BOOT):
        i = rng.integers(0, n, n)
        boots.append(macro_f1(list(p[i]), list(t[i])))
    lo, hi = np.percentile(boots, [2.5, 97.5])
    return val, lo, hi


def mcnemar(base_correct: Dict[int, bool], maa_correct: Dict[int, bool]) -> Tuple[int, int, float]:
    """Return (b, c, p) where b=base-correct&MAA-wrong, c=base-wrong&MAA-correct."""
    ids = set(base_correct) & set(maa_correct)
    b = sum(1 for i in ids if base_correct[i] and not maa_correct[i])
    c = sum(1 for i in ids if not base_correct[i] and maa_correct[i])
    n = b + c
    if n == 0:
        return b, c, 1.0
    # exact binomial two-sided (robust for any n)
    from math import comb
    k = min(b, c)
    p = min(1.0, 2 * sum(comb(n, i) for i in range(0, k + 1)) / (2 ** n))
    return b, c, p


def confusion(pred: List[str], truth: List[str], order: List[str]) -> Tuple[np.ndarray, List[str]]:
    labels = list(order)
    for v in list(pred) + list(truth):
        if v and v not in labels:
            labels.append(v)
    if "" in set(pred):
        labels.append("(empty)")
    lab_i = {l: i for i, l in enumerate(labels)}
    mat = np.zeros((len(labels), len(labels)), dtype=int)
    for p, t in zip(pred, truth):
        if t == "":
            continue
        pp = p if p else "(empty)"
        if pp in lab_i and t in lab_i:
            mat[lab_i[t], lab_i[pp]] += 1
    # trim all-zero rows/cols
    keep = [i for i, l in enumerate(labels) if mat[i].sum() > 0 or mat[:, i].sum() > 0]
    labels = [labels[i] for i in keep]
    mat = mat[np.ix_(keep, keep)]
    return mat, labels


def save_confusion_png(mat, labels, title, path):
    """Render one confusion matrix. Returns False if plotting is unavailable.

    The tables are the deliverable; a host without matplotlib should still produce
    them rather than failing at the figure step. Set NO_FIGURES=1 to skip plotting.
    """
    if os.environ.get("NO_FIGURES"):
        return False
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        return False
    fig, ax = plt.subplots(figsize=(1.1 + 0.5 * len(labels), 1.0 + 0.5 * len(labels)))
    im = ax.imshow(mat, cmap="Blues")
    ax.set_xticks(range(len(labels))); ax.set_xticklabels(labels, rotation=45, ha="right", fontsize=8)
    ax.set_yticks(range(len(labels))); ax.set_yticklabels(labels, fontsize=8)
    ax.set_xlabel("Predicted"); ax.set_ylabel("Ground truth")
    ax.set_title(title, fontsize=9)
    thr = mat.max() / 2 if mat.max() else 0
    for i in range(len(labels)):
        for j in range(len(labels)):
            if mat[i, j]:
                ax.text(j, i, str(mat[i, j]), ha="center", va="center", fontsize=7,
                        color="white" if mat[i, j] > thr else "black")
    fig.tight_layout()
    fig.savefig(path, dpi=150); plt.close(fig)
    return True


# ---- main ----------------------------------------------------------------
def main() -> None:
    gt = load_gt()
    runs = discover()
    print(f"GT rows: {len(gt)} | runs discovered: {len(runs)}")
    if EXCLUDED:
        print(f"  (sensitivity analysis: {len(EXCLUDED)} case(s) excluded from every run)")
    if ONLY_MODELS or ONLY_EDITIONS:
        print(f"  (filtered: MODELS={sorted(ONLY_MODELS) or 'all'}, EDITIONS={sorted(ONLY_EDITIONS) or 'all'})")
    for r in runs:
        short = "  << INCOMPLETE" if len(r["data"]) < len(gt) else ""
        print(f"  - {r['file']}  ({MODEL_LABEL[r['model']]}, {MODE_LABEL[r['mode']]}, "
              f"AJCC {r['ed']}th, n={len(r['data'])}){short}")

    axes = ["T", "N", "M", "S"]
    # ---- Table 4 (accuracy + 95% CI, strict; plus lenient N/M for 9th) ----
    rows_md = []
    with open(os.path.join(OUT, "table4.csv"), "w") as fcsv:
        fcsv.write("model,mode,edition,axis,n,accuracy,ci_lo,ci_hi,lenient\n")
        for r in sorted(runs, key=lambda x: (x["ed"], x["model"], x["mode"])):
            for ax in axes:
                p, t = preds(r, ax), gts(gt, r, ax)
                acc, lo, hi, n = accuracy_ci(p, t, f"{r['ed']}|{r['model']}|{r['mode']}|{ax}")
                len_acc = ""
                if r["ed"] == "9" and ax in ("N", "M"):
                    len_acc = f"{accuracy_lenient(p, t, ax):.1f}"
                fcsv.write(f"{MODEL_LABEL[r['model']]},{MODE_LABEL[r['mode']]},{r['ed']}th,{ax},{n},"
                           f"{acc:.1f},{lo:.1f},{hi:.1f},{len_acc}\n")
                rows_md.append((r["ed"], MODEL_LABEL[r["model"]], MODE_LABEL[r["mode"]], ax,
                                f"{acc:.1f} ({lo:.1f}–{hi:.1f})" + (f" [len {len_acc}]" if len_acc else "")))

    # markdown Table 4 pivoted: rows = model×mode, cols = T/N/M/Stage, per edition
    with open(os.path.join(OUT, "table4.md"), "w") as fmd:
        fmd.write("# Table 4 — accuracy % (bootstrap 95% CI); [len]=9th family-lenient N/M\n")
        for ed in ("8", "9"):
            fmd.write(f"\n## AJCC {ed}th\n\n| Model | Mode | T | N | M | Stage |\n|---|---|---|---|---|---|\n")
            cell = defaultdict(dict)
            for (e, model, mode, ax, txt) in rows_md:
                if e == ed:
                    cell[(model, mode)][ax] = txt
            for (model, mode) in sorted(cell, key=lambda k: (k[0], LABEL_RANK.get(k[1], 99))):
                c = cell[(model, mode)]
                fmd.write(f"| {model} | {mode} | {c.get('T','')} | {c.get('N','')} | {c.get('M','')} | {c.get('S','')} |\n")

    # ---- Macro-F1 (+95% CI) and per-class F1 (AR-1) ----
    f1_rows = defaultdict(dict)
    with open(os.path.join(OUT, "table4_f1.csv"), "w") as ff, \
         open(os.path.join(OUT, "per_class_f1.csv"), "w") as fpc:
        ff.write("model,mode,edition,axis,macro_f1,ci_lo,ci_hi\n")
        fpc.write("model,mode,edition,axis,class,precision,recall,f1,support\n")
        for r in sorted(runs, key=lambda x: (x["ed"], x["model"], x["mode"])):
            for ax in axes:
                p, t = preds(r, ax), gts(gt, r, ax)
                mf, lo, hi = macro_f1_ci(p, t, f"{r['ed']}|{r['model']}|{r['mode']}|{ax}")
                ff.write(f"{MODEL_LABEL[r['model']]},{MODE_LABEL[r['mode']]},{r['ed']}th,{ax},{mf:.1f},{lo:.1f},{hi:.1f}\n")
                f1_rows[(r['ed'], MODEL_LABEL[r['model']], MODE_LABEL[r['mode']])][ax] = f"{mf:.1f} ({lo:.1f}–{hi:.1f})"
                for c, (pr, rc, f1, sup) in _per_class(p, t).items():
                    fpc.write(f"{MODEL_LABEL[r['model']]},{MODE_LABEL[r['mode']]},{r['ed']}th,{ax},{c},{pr:.1f},{rc:.1f},{f1:.1f},{sup}\n")
    with open(os.path.join(OUT, "table4_f1.md"), "w") as fmd:
        fmd.write("# Table 4 — macro-F1 % (bootstrap 95% CI)\n")
        for ed in ("8", "9"):
            fmd.write(f"\n## AJCC {ed}th\n\n| Model | Mode | T | N | M | Stage |\n|---|---|---|---|---|---|\n")
            for (e, model, mode) in sorted(f1_rows, key=lambda k: (k[1], LABEL_RANK.get(k[2], 99))):
                if e != ed:
                    continue
                c = f1_rows[(e, model, mode)]
                fmd.write(f"| {model} | {mode} | {c.get('T','')} | {c.get('N','')} | {c.get('M','')} | {c.get('S','')} |\n")

    # ---- McNemar: every pair of conditions present, per model×edition×axis ----
    def _correct_map(run: dict, ax: str) -> Dict[int, bool]:
        key = ax if ax != "S" else "Stage"
        truth = gts(gt, run, ax)
        return {rec.get("id"): ((rec.get("aiTnm") or {}).get(key, "") == g)
                for rec, g in zip(run["data"], truth) if g != ""}

    with open(os.path.join(OUT, "mcnemar.md"), "w") as fm:
        fm.write("# McNemar — paired comparison of conditions on the identical cases.\n")
        fm.write("\nConditions: baseline (A, single prompt), decomposition-only (C), MAA (D).\n")
        fm.write("b = left-only-correct, c = right-only-correct; n = cases scored in BOTH runs,\n")
        fm.write("so a partially completed run narrows the pair rather than being scored against the full set.\n")
        fm.write("Accuracies below are computed on those same n cases and therefore may differ from Table 4.\n")
        fm.write("\n| Model | Edition | Axis | Comparison | n | left acc | right acc | b | c | p (exact) |\n")
        fm.write("|---|---|---|---|---|---|---|---|---|---|\n")
        by = defaultdict(dict)
        for r in runs:
            by[(r["model"], r["ed"])][r["mode"]] = r
        for (model, ed), modes in sorted(by.items()):
            present = [m for m in MODE_ORDER if m in modes]
            for i in range(len(present)):
                for j in range(i + 1, len(present)):
                    ml, mr = present[i], present[j]
                    for ax in axes:
                        lc, rc_ = _correct_map(modes[ml], ax), _correct_map(modes[mr], ax)
                        ids = set(lc) & set(rc_)
                        if not ids:
                            continue
                        b, c, p = mcnemar(lc, rc_)
                        la = 100 * sum(lc[i_] for i_ in ids) / len(ids)
                        ra = 100 * sum(rc_[i_] for i_ in ids) / len(ids)
                        sig = "*" if p < 0.05 else ""
                        fm.write(f"| {MODEL_LABEL[model]} | {ed}th | {ax} | "
                                 f"{MODE_LABEL[ml]} vs {MODE_LABEL[mr]} | {len(ids)} | "
                                 f"{la:.1f} | {ra:.1f} | {b} | {c} | {p:.3g}{sig} |\n")

    # ---- confusion matrices (CSV + PNG) for every run/axis ----
    figs_ok = False
    for r in runs:
        for ax in axes:
            order = T_ORDER if ax == "T" else N_ORDER[r["ed"]] if ax == "N" else M_ORDER[r["ed"]] if ax == "M" else S_ORDER
            mat, labels = confusion(preds(r, ax), gts(gt, r, ax), order)
            tag = f"{r['model']}_{r['mode']}_ajcc{r['ed']}_{ax}"
            # csv
            with open(os.path.join(CONF, tag + ".csv"), "w") as fc:
                fc.write("gt\\pred," + ",".join(labels) + "\n")
                for i, l in enumerate(labels):
                    fc.write(l + "," + ",".join(str(x) for x in mat[i]) + "\n")
            # png
            figs_ok = save_confusion_png(
                mat, labels,
                f"{MODEL_LABEL[r['model']]} {MODE_LABEL[r['mode']]} AJCC{r['ed']}th — {ax}",
                os.path.join(FIG, tag + ".png")) or figs_ok
    print(f"\nWrote: {OUT}/table4.md, table4_f1.md, table4.csv, mcnemar.md, confusion/*.csv")
    print("Figures: " + (f"{FIG}/*.png" if figs_ok
                         else "skipped (matplotlib unavailable or NO_FIGURES set)"))


if __name__ == "__main__":
    main()
